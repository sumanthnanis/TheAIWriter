import os
from langchain_openai import ChatOpenAI
from langchain_community.tools.tavily_search import TavilySearchResults
from langgraph.prebuilt import chat_agent_executor
from langchain_core.messages import HumanMessage
from langchain.utilities.tavily_search import TavilySearchAPIWrapper

api_key = "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"
search = TavilySearchAPIWrapper(tavily_api_key=api_key)
tavily_tool = TavilySearchResults(api_wrapper=search)

tools = [tavily_tool]



model = ChatOpenAI(model='gpt-3.5-turbo-1106',temperature=1,api_key="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")

app = chat_agent_executor.create_function_calling_executor(model, tools)
template = """
         You are a AI research paper writer assistant. 
         
         Use the provided tools based on the query or input if u needed.
         
         you have access to the tools : \n---\n Taveily Search Tool \n---\n "you should provide accurate data  on the research topic. Atleast 1 pages should be the content" \n---\n
         Rule 1 : Don't generate the links in between sections only generate links in reference section

         Here is the user's problem statement - 
        
        


"""

qe = """Based on the user's input you should write the research paper . If you want to leave a place for images or anything  just say <----- title ----->
        Here is the User's question - """


# Read the lines from the file and concatenate them into a single string
with open('paper.txt','r') as fr:
    problem_statement_lines = fr.readlines()
problem_statement = ''.join(problem_statement_lines)





list1 = ['Introduction','Methodology','Results','Conclusion','References','Appendices']
ert = 'write the ' +'abstract'+ 'for it Giving heading as abstract for it'
# inputs = {"messages": [HumanMessage(content=template +problem_statement+qe+ert)]}
# Concatenate the template, problem_statement, qe, and ert into a single string
inputs = {"messages": [HumanMessage(content=template + problem_statement + qe + ert)]}
wd = app.invoke(inputs)['messages'][-1].content

with open('fr1.txt','w') as f:
    f.write(wd)
f.close()
with open('fr1.txt','r') as f:
    data = f.readlines()
f.close()

for i in list1:
    # ert = 'How to write a ' + i + ' in research papers give me the guidelines , what to mention , what not to mention in 4 lines'
    # inputs = {"messages": [HumanMessage(content=ert)]}
    # wd1 = app.invoke(inputs)['messages'][-1].content
    #"Guidelines to write research paper "+wd1+"Above are the rules and guidelines follow this to generate the answer based on the user's problem statemetn and input "+
    ert =  qe+'write the '+i+' for it'
    inputs = {"messages": [HumanMessage(content=template +problem_statement+ert)]}
    wd = app.invoke(inputs)['messages'][-1].content
    with open('fr1.txt','a') as f:
        f.write(wd)



print(''.join(data))


text = 'The development of e-commerce mobile applications has become increasingly important in providing a platform for selling products, especially for Self Help Groups (SHGs) trained to create unique products. This research paper explores the use of Flutter, Dart, Node.js, and MySQL as the key technologies in the development of an e-commerce mobile application for SHGs. These technologies offer efficient cross-platform development, clean and maintainable code, scalability, and a structured database for product visualization. This paper presents a comprehensive overview of the e-commerce mobile application development process and the technologies involved.'
def split_text_into_lines(text, words_per_line):
    words = text.split()
    lines = [' '.join(words[i:i+words_per_line]) for i in range(0, len(words), words_per_line)]
    return '\n'.join(lines)


words_per_line = 11
formatted_text = split_text_into_lines(text, words_per_line)

with open('output.txt', 'w') as file:
    file.write(formatted_text)

print("Text saved to 'output.txt'")

 
from fpdf import FPDF
 
with open('output.txt','r') as f3:
    data = f3.readlines()
# save FPDF() class into a 
# variable pdf
pdf = FPDF()
 
# Add a page
pdf.add_page()
 
# set style and size of font 
# that you want in the pdf
pdf.set_font("Arial", size = 15)
 
# create a cell

for gt in data:


    pdf.cell(200, 10, txt =gt, 
            ln = 1, align = 'C')


 
# save the pdf with name .pdf
pdf.output("GFG.pdf")   

from docx import Document

def convert_txt_to_docx(input_file, output_file):
    # Read the content from the text file
    with open(input_file, 'r') as file:
        content = file.read()

    # Create a new Word document
    doc = Document()

    # Add the content to the document
    doc.add_paragraph(content)

    # Save the document
    doc.save(output_file)

input_file_path = 'output.txt'
output_file_path = 'output.docx'
convert_txt_to_docx(input_file_path, output_file_path)

print("Text file converted to Word document successfully.")
